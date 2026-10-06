import json
import copy
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path
from unittest.mock import patch
from io import BytesIO

from openpyxl import load_workbook

import recomendar as r

ORIGEM = next(Path(__file__).parent.glob("*.xlsx"))


def plano_teste(dados):
    return r.Plano(resumo=r.resumo_base(dados), recomendacoes=[
        r.Proposta(competencia_id=c["id"], decisao="recomendada", acao_id=c["acoes_permitidas"][0],
                   justificativa="Justificativa artificial para verificar a integração.",
                   responsavel_sugerido="Gerente", prazo_dias_sugerido=30,
                   criterio_conclusao="Critério artificial de avaliação usado apenas no teste.",
                   orientacao_ao_gestor="Orientação artificial de avaliação usada apenas no teste.")
        for c in dados["priorizadas"]])


class TestRecomendador(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        cls.dados = r.analisar(ORIGEM)
        cls.dados["metodologia"] = r.carregar_pbl4(ORIGEM.parent / "PBL4_Projeto_do_Artefato_1.docx")

    def test_contexto_extraido_do_pbl4(self):
        contexto = self.dados["metodologia"]
        trechos = {t["id"]: t["texto"] for t in contexto["trechos"]}
        self.assertIn("ação livre", trechos["UC6.alternativo"])
        self.assertIn("em separado", trechos["UC7.principal"])
        self.assertIn("evidência observável", trechos["RF10"])
        self.assertEqual(len(contexto["sha256"]), 64)

    def test_catalogo_sem_candidatas_nao_impede_diagnostico(self):
        with tempfile.TemporaryDirectory() as tmp:
            arquivo = Path(tmp) / "sem_catalogo.xlsx"
            wb = load_workbook(ORIGEM)
            wb["Catálogo de Ações"].delete_rows(2, 8)
            wb.save(arquivo)
            wb.close()
            dados = r.analisar(arquivo)
            self.assertEqual(len(dados["priorizadas"]), 3)
            self.assertTrue(all(not c["acoes_permitidas"] for c in dados["priorizadas"]))

    def test_abstencao_e_campos_nulos(self):
        p = plano_teste(self.dados)
        abstencao = p.recomendacoes[0].model_dump()
        abstencao.update(decisao="sem_acao_adequada", acao_id=None, responsavel_sugerido=None,
                         prazo_dias_sugerido=None, criterio_conclusao=None)
        p.recomendacoes[0] = r.Proposta.model_validate(abstencao)
        r.validar(p, self.dados)
        abstencao["decisao"] = "dados_insuficientes"
        r.Proposta.model_validate(abstencao)
        abstencao["acao_id"] = self.dados["priorizadas"][0]["acoes_permitidas"][0]
        with self.assertRaises(ValueError):
            r.Proposta.model_validate(abstencao)
        abstencao.update(decisao="recomendada", acao_id=None)
        with self.assertRaises(ValueError):
            r.Proposta.model_validate(abstencao)

    def test_ids_catalogo_estaveis_ao_reordenar(self):
        with tempfile.TemporaryDirectory() as tmp:
            arquivo = Path(tmp) / "reordenada.xlsx"
            wb = load_workbook(ORIGEM)
            ws = wb["Catálogo de Ações"]
            linhas = list(ws.iter_rows(min_row=2, max_row=9, max_col=4, values_only=True))
            for indice, linha in enumerate(reversed(linhas), 2):
                for coluna, valor in enumerate(linha, 1):
                    ws.cell(indice, coluna, valor)
            wb.save(arquivo)
            wb.close()
            d = r.analisar(arquivo)
            self.assertEqual({a["acao"]: a["id"] for a in d["catalogo"]},
                             {a["acao"]: a["id"] for a in self.dados["catalogo"]})

    def test_limite_de_checagem_vem_da_planilha(self):
        d = copy.deepcopy(self.dados)
        for s in d["salvaguardas"]:
            if s["ancora"] == "Diretriz de uso crítico":
                s["observacao"] = "Checar manualmente acima de R$ 5.000."
        prompt = json.dumps(r.mensagens(d), ensure_ascii=False)
        self.assertNotIn("R$ 2.000", prompt)
        self.assertIn("R$ 5.000", prompt)
        self.assertIn("validação bilateral", d["status_ciclo"])

    def test_exporta_abstencao_sem_acao_e_com_rastreabilidade(self):
        p = plano_teste(self.dados)
        v = p.recomendacoes[0].model_dump()
        v.update(decisao="sem_acao_adequada", acao_id=None, responsavel_sugerido=None,
                 prazo_dias_sugerido=None, criterio_conclusao=None)
        p.recomendacoes[0] = r.Proposta.model_validate(v)
        with tempfile.TemporaryDirectory() as tmp:
            pasta = Path(tmp)
            r.exportar(ORIGEM, self.dados, p, pasta, {"modo": "teste"})
            resultado = json.loads((pasta / "analise.json").read_text(encoding="utf-8"))
            self.assertIsNone(resultado["rastreabilidade"][0]["cadeia"]["acao"])
            self.assertIn("UC6.alternativo", resultado["rastreabilidade"][0]["referencias_pbl4"])
            wb = load_workbook(pasta / "MRCC-PV_com_propostas.xlsx")
            self.assertIsNone(wb["Propostas LLM"]["E2"].value)
            self.assertIn("sem_acao_adequada", wb["Propostas LLM"]["N2"].value)
            self.assertIn("R$ 2.000", wb["Propostas LLM"]["O2"].value)
            wb.close()

    def test_resumo_nao_declara_progresso(self):
        p = plano_teste(self.dados)
        p.resumo = "O atendente já demonstrou progresso e o ciclo está validado."
        with self.assertRaises(ValueError):
            r.validar(p, self.dados)

    def test_trilha_manual_oculta_por_padrao(self):
        prompt = json.loads(r.mensagens(self.dados)[1]["content"])
        self.assertNotIn("trilha_existente", prompt["dados"])
        d = copy.deepcopy(self.dados)
        d["usar_trilha_contexto"] = True
        prompt = json.loads(r.mensagens(d)[1]["content"])
        self.assertEqual(prompt["dados"]["trilha_existente"], d["trilha_existente"])

    def test_sem_candidatas_restringe_decisao_no_schema(self):
        d = copy.deepcopy(self.dados)
        for c in d["priorizadas"]:
            c["acoes_permitidas"] = []
        props = r.schema_plano(d)["$defs"]["Proposta"]["properties"]
        self.assertEqual(props["decisao"]["enum"], ["sem_acao_adequada"])
        self.assertEqual(props["acao_id"]["type"], "null")

    def test_prioridades_e_salvaguarda_exemplo(self):
        self.assertEqual([c["score"] for c in self.dados["priorizadas"]], [26, 16, 15])
        self.assertEqual([c["lacuna"] for c in self.dados["priorizadas"]], [2, 1, 1])
        self.assertEqual(len(self.dados["competencias"]), 4)
        self.assertFalse(self.dados["pendencias"])
        s = [s for s in self.dados["salvaguardas"] if s["resposta"] == "Não"]
        self.assertEqual(len(s), 1)
        self.assertIn("checagem", s[0]["ajuste"])

    def test_caches_nao_determinam_prioridade_e_ajuste_ausente_bloqueia_ciclo(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "editada.xlsx"
            wb = load_workbook(ORIGEM)
            wb["M2 - Competências"]["E2"] = "N2 - Usuário autônomo"
            wb["Salvaguarda"]["F4"] = None
            wb.save(p)
            wb.close()
            d = r.analisar(p)
            self.assertEqual([c["score"] for c in d["priorizadas"]], [16, 15])
            self.assertTrue(d["pendencias"])

    def test_rejeita_acao_inventada_e_competencia_sem_lacuna(self):
        p = plano_teste(self.dados)
        p.recomendacoes[0].acao_id = "curso_inexistente"
        with self.assertRaises(ValueError):
            r.validar(p, self.dados)
        p = plano_teste(self.dados)
        p.recomendacoes[0].competencia_id = "C4"
        with self.assertRaises(ValueError):
            r.validar(p, self.dados)

    def test_rejeita_duplicacao_e_ordem_errada(self):
        p = plano_teste(self.dados)
        p.recomendacoes = list(reversed(p.recomendacoes))
        with self.assertRaises(ValueError):
            r.validar(p, self.dados)
        p.recomendacoes = [p.recomendacoes[0]] * 3
        with self.assertRaises(ValueError):
            r.validar(p, self.dados)

    def test_preserva_xml_original_e_neutraliza_formula_llm(self):
        with tempfile.TemporaryDirectory() as tmp:
            pasta = Path(tmp)
            p = plano_teste(self.dados)
            p.recomendacoes[0].responsavel_sugerido = '=HYPERLINK("http://example.com")'
            r.exportar(ORIGEM, self.dados, p, pasta, {"modo": "teste"})
            destino = pasta / "MRCC-PV_com_propostas.xlsx"
            with zipfile.ZipFile(ORIGEM) as a, zipfile.ZipFile(destino) as b:
                for name in a.namelist():
                    if name not in ("xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                        self.assertEqual(a.read(name), b.read(name), name)
            wb = load_workbook(destino)
            self.assertEqual(len(wb.sheetnames), 9)
            cell = wb["Propostas LLM"]["H2"]
            self.assertNotEqual(cell.data_type, "f")
            self.assertTrue(cell.value.startswith("'="))
            wb.close()

    def test_api_local_e_validacao_com_resposta_simulada(self):
        p = plano_teste(self.dados)
        raw = {"done": True, "message": {"content": p.model_dump_json()}, "model": r.MODELO}
        with tempfile.TemporaryDirectory() as tmp:
            responses = [BytesIO(b'{"version":"teste"}'), BytesIO(b'{"models":[]}'), BytesIO(json.dumps(raw).encode())]
            with patch("urllib.request.urlopen", side_effect=responses) as api:
                plano, meta = r.gerar(self.dados, r.MODELO, Path(tmp))
                self.assertEqual(len(plano.recomendacoes), 3)
                req = api.call_args.args[0]
                self.assertEqual(req.full_url, r.API)
                self.assertFalse(json.loads(req.data)["stream"])
                self.assertEqual(meta["tentativas"], 1)

    def test_repeticao_corrige_acao_invalida(self):
        p = plano_teste(self.dados)
        valido = {"done": True, "message": {"content": p.model_dump_json()}}
        p.recomendacoes[0].acao_id = "inventada"
        invalido = {"done": True, "message": {"content": p.model_dump_json()}}
        responses = [BytesIO(b'{}'), BytesIO(b'{}'),
                     BytesIO(json.dumps(invalido).encode()), BytesIO(json.dumps(valido).encode())]
        with tempfile.TemporaryDirectory() as tmp:
            with patch("urllib.request.urlopen", side_effect=responses):
                plano, meta = r.gerar(self.dados, r.MODELO, Path(tmp))
                self.assertEqual(meta["tentativas"], 2)
                self.assertEqual(plano.recomendacoes[0].acao_id, self.dados["priorizadas"][0]["acoes_permitidas"][0])

    def test_nivel_invalido(self):
        for valor in (None, "", "N4", "N20", "assistido"):
            with self.assertRaises(ValueError):
                r.nivel(valor)


if __name__ == "__main__":
    unittest.main()
