"""Adaptador da interface web para o mesmo motor de recomendação já validado."""
import argparse
import hashlib
import json
import os
import sys
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
import recomendar as r
import web_llm

ROOT = Path(__file__).parent
TEMPLATE = Path(os.environ.get("MRCC_TEMPLATE", ROOT / "MRCC-PV_Exemplo_Preenchido_1 (1).xlsx"))
PBL4 = ROOT / "PBL4_Projeto_do_Artefato_1.docx"
CAMPOS = {
    "tarefas": ("M1 - Diagnóstico", ["funcao", "tarefa", "frequencia", "criticidade", "ferramenta", "efeito_ia"]),
    "competencias": ("M2 - Competências", ["funcao", "competencia", "eixo", "tarefa", "atual", "alvo"]),
    "catalogo": ("Catálogo de Ações", ["eixo", "acao", "recurso", "custo"]),
    "salvaguardas": ("Salvaguarda", ["numero", "ancora", "verificacao", "resposta", "observacao", "ajuste"]),
    "trilha": ("M4 - Trilha", ["competencia", "prioridade", "acao", "recurso", "responsavel", "prazo_dias", "criterio"]),
}


def entrada_da_planilha(caminho):
    with zipfile.ZipFile(caminho) as z:
        if sum(i.file_size for i in z.infolist()) > 25_000_000:
            raise ValueError("A planilha é grande demais para este protótipo.")
    dados = r.analisar(caminho)
    entrada = {nome: [{campo: linha.get(campo) for campo in campos} for linha in dados["trilha_existente" if nome == "trilha" else nome]] for nome, (_, campos) in CAMPOS.items()}
    wb = load_workbook(caminho, read_only=True)
    try:
        entrada["pesos_criticidade"] = {r.texto(wb["Legenda"].cell(i, 5).value): wb["Legenda"].cell(i, 6).value for i in range(3, 6)}
        entrada["pesos_frequencia"] = {r.texto(wb["Legenda"].cell(i, 8).value): wb["Legenda"].cell(i, 9).value for i in range(3, 6)}
    finally:
        wb.close()
    return entrada


def validar_entrada(entrada):
    prazo = entrada.get('prazo_dias_proposto', 14)
    if type(prazo) is not int or not 1 <= prazo <= 180:
        raise ValueError('Informe um prazo inicial inteiro entre 1 e 180 dias.')
    limites = {"tarefas": 12, "competencias": 8, "catalogo": 16, "salvaguardas": 6, "trilha": 12}
    for nome, (_, campos) in CAMPOS.items():
        linhas = entrada.get(nome, [])
        if not isinstance(linhas, list) or len(linhas) > limites[nome]:
            raise ValueError(f"Quantidade inválida de linhas em {nome} (máximo {limites[nome]}).")
        if nome in ("tarefas", "competencias", "salvaguardas") and not linhas:
            raise ValueError(f"Preencha {nome} antes de continuar.")
        for linha in linhas:
            if not isinstance(linha, dict):
                raise ValueError(f"Linha inválida em {nome}.")
            for campo in campos:
                valor = linha.get(campo)
                if valor is not None and (type(valor) not in (str, int, float) or len(str(valor)) > 1000):
                    raise ValueError(f"Conteúdo inválido em {nome}/{campo}.")
    if len(entrada["salvaguardas"]) != 6:
        raise ValueError("A planilha deve ter as seis âncoras de Salvaguarda.")
    for nome, labels in (("pesos_criticidade", ["Alta", "Média", "Baixa"]), ("pesos_frequencia", ["Diária", "Semanal", "Eventual"])):
        pesos = entrada.get(nome, dict(zip(labels, [3, 2, 1])))
        if not isinstance(pesos, dict) or set(pesos) != set(labels) or any(type(v) is not int or v < 1 or v > 100 for v in pesos.values()):
            raise ValueError("Pesos de frequência/criticidade inválidos.")


def escrever_entrada(entrada, destino):
    validar_entrada(entrada)
    wb = load_workbook(TEMPLATE)
    try:
        for nome, col in (("pesos_criticidade", 5), ("pesos_frequencia", 8)):
            if nome in entrada:
                for i in range(3, 6):
                    wb["Legenda"].cell(i, col + 1, entrada[nome][r.texto(wb["Legenda"].cell(i, col).value)])
        for nome, (aba, campos) in CAMPOS.items():
            ws = wb[aba]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for linha in entrada.get(nome, []):
                ws.append([r.celula_segura(linha.get(campo)) for campo in campos])
        wb.save(destino)
    finally:
        wb.close()
    dados = r.analisar(destino)
    # O Excel exportado apresenta M3 atualizado mesmo sem recálculo pelo Excel.
    wb = load_workbook(destino)
    try:
        ws = wb["M3 - Lacunas"]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for c in dados["competencias"]:
            t = c["contexto_tarefa"]
            pc = next(wb["Legenda"].cell(i, 6).value for i in range(3, 6) if wb["Legenda"].cell(i, 5).value == t["criticidade"])
            pf = next(wb["Legenda"].cell(i, 9).value for i in range(3, 6) if wb["Legenda"].cell(i, 8).value == t["frequencia"])
            ws.append([r.celula_segura(v) for v in [c["competencia"], c["tarefa"], c["atual"], c["alvo"], c["atual_num"], c["alvo_num"], c["lacuna"], t["criticidade"], t["frequencia"], pc, pf, c["score"], c["prioridade"]]])
        wb.save(destino)
    finally:
        wb.close()
    return dados


def processar(entrada, pasta, modo):
    pasta.mkdir(parents=True, exist_ok=True)
    caminho = pasta / "entrada.xlsx"
    dados = escrever_entrada(entrada, caminho)
    dados["metodologia"] = r.carregar_pbl4(PBL4)
    dados["usar_trilha_contexto"] = False
    dados["prompt_compacto"] = True
    dados['parametros_gestor'] = {'prazo_dias_proposto': entrada.get('prazo_dias_proposto', 14)}
    if len(json.dumps(r.mensagens(dados), ensure_ascii=False)) > 28000:
        raise ValueError("O formulário está extenso demais. Reduza o número de competências ou o tamanho das descrições.")
    meta = {"versao": r.VERSAO, "modo": modo, "origem": "formulario_web",
            "gerado_em_utc": datetime.now(timezone.utc).isoformat(), "modelo": r.MODELO if modo == "llm" else None,
            "versao_prompt": web_llm.VERSAO, "trilha_manual_enviada_a_llm": False,
            "origem_criterios": "referencias explicitas do prototipo; tarefa e nivel-alvo, com revisao humana pendente",
            "sha256_planilha": hashlib.sha256(caminho.read_bytes()).hexdigest(),
            "sha256_pbl4": dados["metodologia"]["sha256"],
            "formula_score": "max(0, alvo-atual)*10 + peso_criticidade + peso_frequencia"}
    r.exportar(caminho, dados, None, pasta, meta)
    if modo == "llm":
        if dados["priorizadas"]:
            plano, metricas = web_llm.gerar(dados, r.MODELO, pasta)
            meta.update(metricas)
        else:
            plano = r.Plano(resumo=r.resumo_base(dados), recomendacoes=[])
        r.exportar(caminho, dados, plano, pasta, meta)
    return json.loads((pasta / "analise.json").read_text(encoding="utf-8"))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("acao", choices=["seed", "importar", "analise", "llm"])
    parser.add_argument("--entrada", type=Path)
    parser.add_argument("--saida", type=Path)
    args = parser.parse_args()
    if args.acao in ("seed", "importar"):
        entrada = entrada_da_planilha(TEMPLATE if args.acao == "seed" else args.entrada)
        validar_entrada(entrada)
        print(json.dumps(entrada, ensure_ascii=False))
    else:
        resultado = processar(json.loads(args.entrada.read_text(encoding="utf-8")), args.saida, args.acao)
        print(json.dumps({"sucesso": True, "priorizadas": len(resultado["diagnostico"]["priorizadas"])}))


if __name__ == "__main__":
    warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        main()
    except Exception as e:
        print(str(e), file=sys.stderr)
        sys.exit(1)
